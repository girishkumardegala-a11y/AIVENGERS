"""
extractor.py
------------
Handles raw text extraction from uploaded Job Description files (PDF / DOCX).
Strips repeated headers/footers/page numbers and collapses excess whitespace.
"""

import io
import re
from collections import Counter

import pdfplumber
from docx import Document


class UnsupportedFileType(Exception):
    pass


def extract_text(filename: str, file_bytes: bytes) -> str:
    """Route to the correct extractor based on file extension."""
    lower = filename.lower()
    if lower.endswith(".pdf"):
        raw = _extract_pdf(file_bytes)
    elif lower.endswith(".docx"):
        raw = _extract_docx(file_bytes)
    else:
        raise UnsupportedFileType(
            f"Unsupported file type: {filename}. Only .pdf and .docx are supported."
        )
    return _clean_text(raw)


def _extract_pdf(file_bytes: bytes) -> str:
    pages_text = []
    with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
        for page in pdf.pages:
            text = page.extract_text() or ""
            pages_text.append(text)

    # Detect lines repeated on nearly every page (headers/footers) and drop them
    line_counts = Counter()
    per_page_lines = []
    for text in pages_text:
        lines = [l.strip() for l in text.split("\n") if l.strip()]
        per_page_lines.append(lines)
        for l in set(lines):
            line_counts[l] += 1

    n_pages = max(len(pages_text), 1)
    repeated = {
        line for line, count in line_counts.items()
        if n_pages > 1 and count >= max(2, int(n_pages * 0.6))
    }

    cleaned_pages = []
    for lines in per_page_lines:
        kept = [l for l in lines if l not in repeated and not _looks_like_page_number(l)]
        cleaned_pages.append("\n".join(kept))

    return "\n\n".join(cleaned_pages)


def _extract_docx(file_bytes: bytes) -> str:
    doc = Document(io.BytesIO(file_bytes))
    parts = []

    for para in doc.paragraphs:
        if para.text.strip():
            parts.append(para.text.strip())

    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))

    return "\n".join(parts)


def _looks_like_page_number(line: str) -> bool:
    return bool(re.fullmatch(r"(page\s*)?\d+(\s*/\s*\d+)?", line.strip(), re.IGNORECASE))


def _clean_text(text: str) -> str:
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()
